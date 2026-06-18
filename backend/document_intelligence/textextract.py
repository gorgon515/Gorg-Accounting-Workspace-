"""Real text extraction from documents — PDF (text layer), Excel, Word, email,
CSV, and plain text. No fabrication: each method extracts actual content, and
reports a confidence of 0 (with needs_ocr=True) when a PDF has no text layer.
"""
from __future__ import annotations

import csv
import io
import os
from email import message_from_string, policy
from typing import Optional


def extract_pdf(path: str) -> dict:
    from pypdf import PdfReader
    reader = PdfReader(path)
    pages = []
    for p in reader.pages:
        pages.append(p.extract_text() or "")
    text = "\n".join(pages).strip()
    return {"format": "pdf", "method": "pdf_text_layer", "text": text, "pages": len(pages),
            "confidence": 1.0 if text else 0.0, "needs_ocr": not bool(text)}


def extract_xlsx(path: str) -> dict:
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    lines = []
    for ws in wb.worksheets:
        lines.append(f"# Sheet: {ws.title}")
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                lines.append("\t".join(cells))
    return {"format": "xlsx", "method": "openpyxl", "text": "\n".join(lines), "confidence": 1.0}


def extract_docx(path: str) -> dict:
    import docx
    doc = docx.Document(path)
    text = "\n".join(p.text for p in doc.paragraphs if p.text)
    for table in doc.tables:
        for row in table.rows:
            text += "\n" + "\t".join(c.text for c in row.cells)
    return {"format": "docx", "method": "python-docx", "text": text.strip(), "confidence": 1.0}


def extract_eml(path: str) -> dict:
    with open(path, "r", encoding="utf-8", errors="replace") as f:
        msg = message_from_string(f.read(), policy=policy.default)
    body = ""
    if msg.is_multipart():
        for part in msg.walk():
            if part.get_content_type() == "text/plain":
                body += part.get_content()
    else:
        body = msg.get_content()
    header = f"From: {msg['from']}\nTo: {msg['to']}\nSubject: {msg['subject']}\nDate: {msg['date']}\n"
    return {"format": "eml", "method": "email", "text": header + "\n" + (body or ""), "confidence": 1.0,
            "headers": {"from": msg["from"], "subject": msg["subject"], "date": msg["date"]}}


def extract_csv(path: str) -> dict:
    with open(path, "r", encoding="utf-8", errors="replace") as f:
        rows = list(csv.reader(f))
    return {"format": "csv", "method": "csv", "text": "\n".join("\t".join(r) for r in rows), "confidence": 1.0}


def extract_txt(path: str) -> dict:
    with open(path, "r", encoding="utf-8", errors="replace") as f:
        return {"format": "txt", "method": "plain", "text": f.read(), "confidence": 1.0}


_IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".gif"}
_DISPATCH = {".pdf": extract_pdf, ".xlsx": extract_xlsx, ".xlsm": extract_xlsx, ".docx": extract_docx,
             ".eml": extract_eml, ".csv": extract_csv, ".txt": extract_txt, ".md": extract_txt}


def extract_text(path: str) -> dict:
    """Extract real text from a document. Images route to OCR (gated)."""
    ext = os.path.splitext(path)[1].lower()
    if ext in _IMAGE_EXTS:
        from . import ocr
        return ocr.ocr_image(path)
    fn = _DISPATCH.get(ext)
    if not fn:
        raise ValueError(f"unsupported document type: {ext or '(none)'}")
    return fn(path)


def process_bytes(filename: str, data: bytes) -> dict:
    """Write bytes to a temp file (preserving extension) and extract."""
    import tempfile
    ext = os.path.splitext(filename)[1].lower() or ".bin"
    with tempfile.NamedTemporaryFile(suffix=ext, delete=False) as tmp:
        tmp.write(data)
        tmp_path = tmp.name
    try:
        out = extract_text(tmp_path)
        out["filename"] = filename
        return out
    finally:
        try:
            os.unlink(tmp_path)
        except OSError:
            pass
